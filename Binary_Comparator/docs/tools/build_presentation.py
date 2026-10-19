"""Build docs/Firmware_Binary_Comparison_Tool_Demo_v1.0.pptx (python-pptx) with screenshots and speaker notes."""

from __future__ import annotations

import os

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

DOCS = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
IMG = os.path.join(DOCS, "images")
OUT = os.path.join(DOCS, "Firmware_Binary_Comparison_Tool_Demo_v1.0.pptx")

SLATE = RGBColor(0x1E, 0x29, 0x3B)
INK = RGBColor(0x1F, 0x29, 0x33)
MUTED = RGBColor(0x52, 0x60, 0x6D)
BLUE = RGBColor(0x25, 0x63, 0xEB)
PURPLE = RGBColor(0x7C, 0x3A, 0xED)
RED = RGBColor(0xC0, 0x26, 0x26)
GREEN = RGBColor(0x15, 0x80, 0x3D)
AMBER = RGBColor(0xA1, 0x62, 0x07)
LIGHT = RGBColor(0xF4, 0xF6, 0xF9)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Segoe UI"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]
W, H = prs.slide_width, prs.slide_height
_page = [0]


# ------------------------------------------------------------------ helpers
def text(slide, left, top, width, height, value, size=18, bold=False, color=INK, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    lines = value if isinstance(value, list) else [value]
    for i, line in enumerate(lines):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align
        run = para.add_run()
        run.text = line
        run.font.name = FONT
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
    return box


def bullets(slide, left, top, width, height, items, size=17, color=INK, space=8):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.space_after = Pt(space)
        head, body = item if isinstance(item, tuple) else ("", item)
        r0 = para.add_run()
        r0.text = "▪  "
        r0.font.color.rgb = BLUE
        r0.font.size = Pt(size)
        r0.font.name = FONT
        if head:
            r = para.add_run()
            r.text = head
            r.font.bold = True
            r.font.size = Pt(size)
            r.font.name = FONT
            r.font.color.rgb = color
        r = para.add_run()
        r.text = body
        r.font.size = Pt(size)
        r.font.name = FONT
        r.font.color.rgb = color
    return box


def rect(slide, left, top, width, height, fill, line=None, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, left, top, width, height)
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(1.25)
    s.shadow.inherit = False
    return s


def picture(slide, name, left, top, max_w, max_h, border=True):
    path = os.path.join(IMG, name)
    with Image.open(path) as im:
        iw, ih = im.size
    scale = min(max_w / iw, max_h / ih)
    w, h = int(iw * scale), int(ih * scale)
    x = left + (max_w - w) // 2
    y = top + (max_h - h) // 2
    pic = slide.shapes.add_picture(path, Emu(x), Emu(y), Emu(w), Emu(h))
    if border:
        pic.line.color.rgb = RGBColor(0xC3, 0xCC, 0xD6)
        pic.line.width = Pt(0.75)
    return pic


def callout(slide, left, top, width, value, color=BLUE, size=13):
    box = rect(slide, left, top, width, Inches(0.55), color, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    box.adjustments[0] = 0.25
    tf = box.text_frame
    tf.margin_left = tf.margin_right = Inches(0.12)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    para = tf.paragraphs[0]
    para.alignment = PP_ALIGN.CENTER
    r = para.add_run()
    r.text = value
    r.font.name = FONT
    r.font.size = Pt(size)
    r.font.bold = True
    r.font.color.rgb = WHITE
    return box


def notes(slide, value):
    slide.notes_slide.notes_text_frame.text = value


def content_slide(title, subtitle=None, section=None):
    s = prs.slides.add_slide(BLANK)
    _page[0] += 1
    rect(s, 0, 0, W, H, LIGHT)
    rect(s, 0, 0, W, Inches(1.05), SLATE)
    rect(s, 0, Inches(1.05), W, Inches(0.06), BLUE)
    text(s, Inches(0.55), Inches(0.16), Inches(10.5), Inches(0.6), title, size=28, bold=True, color=WHITE)
    if subtitle:
        text(s, Inches(0.55), Inches(0.62), Inches(11.5), Inches(0.4), subtitle, size=14,
             color=RGBColor(0xCB, 0xD5, 0xE1))
    if section:
        text(s, Inches(10.3), Inches(0.22), Inches(2.6), Inches(0.4), section.upper(), size=11, bold=True,
             color=RGBColor(0x93, 0xC5, 0xFD), align=PP_ALIGN.RIGHT)
    text(s, Inches(0.55), Inches(7.08), Inches(8), Inches(0.3), "Firmware Binary Comparison Tool  ·  v1.0.0",
         size=10, color=MUTED)
    text(s, Inches(11.8), Inches(7.08), Inches(1), Inches(0.3), str(_page[0]), size=10, color=MUTED,
         align=PP_ALIGN.RIGHT)
    return s


def tile(slide, left, top, width, height, title_, body, color):
    rect(slide, left, top, width, height, WHITE, line=RGBColor(0xD8, 0xDE, 0xE6))
    rect(slide, left, top, Inches(0.09), height, color)
    text(slide, left + Inches(0.25), top + Inches(0.15), width - Inches(0.35), Inches(0.45), title_, size=17,
         bold=True, color=color)
    text(slide, left + Inches(0.25), top + Inches(0.62), width - Inches(0.35), height - Inches(0.7), body, size=13,
         color=MUTED)


AREA_TOP = Inches(1.35)
AREA_H = Inches(5.6)


# ------------------------------------------------------------------ slides
def s_title():
    s = prs.slides.add_slide(BLANK)
    _page[0] += 1
    rect(s, 0, 0, W, H, SLATE)
    rect(s, 0, Inches(5.05), W, Inches(0.08), BLUE)
    # decorative memory-map grid
    for r in range(6):
        for c in range(10):
            color = RED if (r, c) in ((1, 6), (2, 3), (2, 4), (4, 8)) else (
                RGBColor(0x33, 0x41, 0x55) if (r + c) % 3 else RGBColor(0x3B, 0x4A, 0x60))
            rect(s, Inches(8.6 + c * 0.42), Inches(0.9 + r * 0.42), Inches(0.36), Inches(0.36), color)
    text(s, Inches(0.8), Inches(1.6), Inches(7.6), Inches(1.0), "Firmware Binary", size=44, bold=True, color=WHITE)
    text(s, Inches(0.8), Inches(2.4), Inches(7.6), Inches(1.0), "Comparison Tool", size=44, bold=True, color=WHITE)
    text(s, Inches(0.8), Inches(3.45), Inches(7.6), Inches(1.0),
         "Know exactly what changed between two firmware builds —\nwhere, in which function, and whether it was "
         "expected.", size=18, color=RGBColor(0xCB, 0xD5, 0xE1))
    text(s, Inches(0.8), Inches(5.45), Inches(11), Inches(0.5), "Demo  ·  v1.0.0 frozen baseline  ·  16 October 2026",
         size=16, bold=True, color=RGBColor(0x93, 0xC5, 0xFD))
    text(s, Inches(0.8), Inches(5.95), Inches(11), Inches(0.5),
         ".bin  ·  Intel .hex  ·  .elf     |     GUI  ·  CLI  ·  HTML / CSV / JSON reports", size=14,
         color=RGBColor(0x94, 0xA3, 0xB8))
    notes(s, "Introduce the tool: a desktop and command-line utility for embedded firmware engineers. "
             "Goal of the demo: show how quickly we can answer 'what exactly changed between these two builds?'.")


def s_problem():
    s = content_slide("Why a byte diff is not enough", "Questions engineers actually ask before releasing firmware",
                      "Motivation")
    cards = [
        ("Where did it change?", "MCU addresses, not file offsets. BIN, HEX and ELF use different address views.", BLUE),
        ("Was it expected?", "Version strings, build timestamps and CRCs always change — real changes must stand out.",
         AMBER),
        ("Which function?", "Removing one call shifts all later code: hundreds of functions differ without being "
                            "edited.", RED),
        ("Which controller?", "Is this image really for the STM32F411 / PIC32MX575 we expect?", PURPLE),
    ]
    cw = Inches(2.95)
    for i, (t, b, c) in enumerate(cards):
        tile(s, Inches(0.55) + i * (cw + Inches(0.17)), Inches(1.55), cw, Inches(2.2), t, b, c)
    bullets(s, Inches(0.6), Inches(4.15), Inches(12), Inches(2.6), [
        ("Generic diff tools ", "show offsets and raw bytes — no regions, no expectations, no symbols."),
        ("Manual review ", "of hex dumps is slow and misses changes hidden behind code shifts."),
        ("This tool ", "answers all four questions in one view, and as a CI gate from the command line."),
    ], size=18)
    notes(s, "Set the scene with the four questions. Mention the watchdog example we will use later: a 3-line change "
             "produced 92,000 different bytes.")


def s_features():
    s = content_slide("What the tool does", "Feature overview", "Overview")
    feats = [
        ("Byte comparison", "BIN / HEX / ELF by MCU address, size changes, gap fill, 64 MB in ~0.1 s", BLUE),
        ("Regions & expectations", "Named regions, expected areas (version, CRC …), EXPECTED vs UNEXPECTED verdict",
         AMBER),
        ("Functions & modules", "ELF / linker map symbols, relocation-aware: real edits vs side effects", RED),
        ("Controller identification", "STM32F411RET, PIC32MX575F512L … with evidence and confidence", PURPLE),
        ("Hex viewer & memory map", "Synchronised hex view, navigation, search, filters, visual block map", GREEN),
        ("Reports & CLI", "HTML, CSV, JSON; firmware-diff with CI exit codes", SLATE),
    ]
    cw, ch = Inches(3.95), Inches(2.45)
    for i, (t, b, c) in enumerate(feats):
        tile(s, Inches(0.55) + (i % 3) * (cw + Inches(0.2)), Inches(1.5) + (i // 3) * (ch + Inches(0.25)), cw, ch,
             t, b, c)
    notes(s, "Six capability areas. Everything runs locally, needs only Python 3.11, and never modifies the firmware "
             "files.")


def s_architecture():
    s = content_slide("Architecture", "GUI-independent engine shared by GUI, CLI and tests", "Design")
    picture(s, "00_architecture.png", Inches(0.5), AREA_TOP, Inches(12.3), AREA_H, border=False)
    notes(s, "Layers: formats parse and validate files; core holds the engine (IntervalSet-based differences, "
             "analysis, symbols, controller detection); session orchestrates; GUI, CLI and reports are thin front "
             "ends. Standard library only.")


def s_pipeline():
    s = content_slide("How a comparison runs", "Five steps — the expensive one runs once", "Design")
    picture(s, "00_pipeline.png", Inches(0.5), Inches(1.4), Inches(12.3), Inches(2.9), border=False)
    bullets(s, Inches(0.7), Inches(4.45), Inches(12), Inches(2.4), [
        ("Union of addresses: ", "bytes present in only one image are differences — nothing is silently ignored."),
        ("Hierarchical scan: ", "64 KiB chunks, identical areas skipped at memcmp speed, differences stored as runs."),
        ("Instant re-analysis: ", "editing regions, expected areas or symbol files does not re-compare bytes."),
    ], size=17)
    notes(s, "Explain the union-of-addresses idea and the run-based storage that keeps large images fast.")


def s_files():
    s = content_slide("Step 1 · Select the firmware", "Drop zones, base address, symbol files, configuration",
                      "Demo")
    picture(s, "02_files_loaded.png", Inches(0.45), AREA_TOP, Inches(9.2), AREA_H)
    bullets(s, Inches(9.85), Inches(1.6), Inches(3.3), Inches(5.3), [
        ("A / B cards: ", "click or drop a file; name, size, SHA-256, format, address range"),
        ("Base address ", "for raw .bin files"),
        ("Controller ", "identified automatically"),
        ("Symbol files ", "(ELF / .map) unlock function-level results"),
        ("Compare ", "runs automatically → Summary"),
    ], size=14, space=10)
    notes(s, "Live: drop Analysis_Board_Firmware.hex (old) and the watchdog-commented build (new), load both ELFs as "
             "symbol files. Point out the controller row: STM32F411RET, Cortex-M4 with FPU.")


def s_summary():
    s = content_slide("Step 2 · The answer at a glance", "Verdict, key numbers and what changed", "Demo")
    picture(s, "03_summary_watchdog.png", Inches(0.45), AREA_TOP, Inches(9.2), AREA_H)
    callout(s, Inches(9.9), Inches(1.7), Inches(3.1), "92,007 bytes differ (76 %)", RED)
    text(s, Inches(9.9), Inches(2.35), Inches(3.1), Inches(1.0), "…because one call was removed and all later code "
                                                                 "shifted.", size=13, color=MUTED)
    callout(s, Inches(9.9), Inches(3.5), Inches(3.1), "2 functions modified", BLUE)
    callout(s, Inches(9.9), Inches(4.2), Inches(3.1), "6 symbols removed", PURPLE)
    text(s, Inches(9.9), Inches(4.9), Inches(3.1), Inches(1.6), "peripheralInit and main lost their watchdog calls; "
                                                                "the IWDG driver was dropped by the linker.", size=13,
         color=MUTED)
    notes(s, "Key message: the byte count is huge, the real change is tiny. The 'What changed' list shows exactly the "
             "watchdog edit.")


def s_functions():
    s = content_slide("Real edits vs side effects", "Relocation-aware function comparison", "Highlight")
    picture(s, "04_changed_functions.png", Inches(0.45), AREA_TOP, Inches(8.4), AREA_H)
    rect(s, Inches(9.1), Inches(1.55), Inches(3.85), Inches(2.25), WHITE, line=RGBColor(0xD8, 0xDE, 0xE6))
    text(s, Inches(9.3), Inches(1.65), Inches(3.5), Inches(0.4), "Before", size=14, bold=True, color=MUTED)
    text(s, Inches(9.3), Inches(2.0), Inches(3.5), Inches(0.7), "350 'modified'", size=30, bold=True, color=RED)
    text(s, Inches(9.3), Inches(2.75), Inches(3.5), Inches(0.4), "After", size=14, bold=True, color=MUTED)
    text(s, Inches(9.3), Inches(3.05), Inches(3.5), Inches(0.7), "2 modified · 6 removed", size=22, bold=True,
         color=GREEN)
    bullets(s, Inches(9.1), Inches(4.05), Inches(3.9), Inches(2.9), [
        "Calls (BL / JAL) and pointers are decoded in both builds",
        "Same target symbol → RELOCATED, not an edit",
        "762 side effects listed separately",
        "ARM Thumb-2 and MIPS32 (PIC32)",
    ], size=13, space=6)
    notes(s, "This is the key innovation. Removing a call shifts later code, changing branch offsets and literal "
             "pools everywhere. The tool resolves each reference to symbol+offset in both builds and only counts real "
             "content changes.")


def s_relocated():
    s = content_slide("Side effects stay visible", "Relocated / moved symbols on their own filter", "Highlight")
    picture(s, "04b_relocated.png", Inches(0.45), AREA_TOP, Inches(9.2), AREA_H)
    bullets(s, Inches(9.85), Inches(1.6), Inches(3.3), Inches(5.3), [
        ("RELOCATED: ", "only call targets / pointers changed"),
        ("MOVED: ", "identical code at a new address"),
        ("LIKELY RELOCATED: ", "references to unnamed data — listed for review, never hidden"),
        ("Module view ", "from linker map files (motor.o, uart.o …)"),
    ], size=14, space=10)
    notes(s, "Nothing is hidden: switch the filter to see the 762 side effects and their notes.")


def s_hex():
    s = content_slide("Hex comparison viewer", "Synchronised, virtualised — scrolls MB-sized images instantly", "Demo")
    picture(s, "05_hex_viewer.png", Inches(0.45), AREA_TOP, Inches(8.6), AREA_H)
    picture(s, "05b_hex_table.png", Inches(9.25), Inches(1.45), Inches(3.75), Inches(2.6))
    bullets(s, Inches(9.25), Inches(4.25), Inches(3.8), Inches(2.7), [
        "First / prev / next / last difference",
        "Go to address or file offset",
        "Hex search in A, B or both",
        "Filters: diffs only, region, block size",
        "Status bar: offsets, region, block, function",
    ], size=13, space=5)
    notes(s, "Double-click peripheralInit in the function list to jump here. Show n / p keys and the 1-byte table "
             "layout (Address | A | B | Status).")


def s_regions():
    s = content_slide("Regions and expected changes", "Version, timestamp and CRC are allowed to change — nothing "
                                                       "else is", "Demo")
    picture(s, "06_summary_expected.png", Inches(0.45), Inches(1.4), Inches(6.2), Inches(4.3))
    picture(s, "07_regions.png", Inches(6.85), Inches(1.4), Inches(6.05), Inches(4.3))
    text(s, Inches(0.45), Inches(5.85), Inches(6.2), Inches(1.0), "Summary: 22 expected bytes (version, timestamp, "
         "git hash, CRC) are separated from 14,179 unexpected ones.", size=14, color=MUTED, align=PP_ALIGN.CENTER)
    text(s, Inches(6.85), Inches(5.85), Inches(6.05), Inches(1.0), "Regions: Bootloader and Calibration show "
         "EXPECTED CHANGE only; Application and Diagnostics really changed.", size=14, color=MUTED,
         align=PP_ALIGN.CENTER)
    notes(s, "Example project with Bootloader / Application / Calibration / Diagnostics and expected areas for "
             "version, build timestamp, git hash and CRC. Bootloader and Calibration are EXPECTED CHANGE only; "
             "Application and Diagnostics have unexpected changes.")


def s_blocks_map():
    s = content_slide("Changed blocks and memory map", "Where the changes are — as a table and as a picture", "Demo")
    picture(s, "08_changed_blocks.png", Inches(0.45), Inches(1.4), Inches(6.15), Inches(4.3))
    picture(s, "09_memory_map.png", Inches(6.75), Inches(1.4), Inches(6.15), Inches(4.3))
    text(s, Inches(0.45), Inches(5.85), Inches(6.15), Inches(1.0), "Every block with size, class, region, file "
         "offsets and the function it falls in; sortable and filterable.", size=14, color=MUTED,
         align=PP_ALIGN.CENTER)
    text(s, Inches(6.75), Inches(5.85), Inches(6.15), Inches(1.0), "Memory map: red = unexpected, amber = expected "
         "only. Click any cell to open the hex viewer there.", size=14, color=MUTED, align=PP_ALIGN.CENTER)
    notes(s, "Blocks can be sorted by address, size or region and filtered by class or size; each row names the "
             "function. The map: red = unexpected, amber = expected only; click a cell to open the hex viewer.")


def s_config():
    s = content_slide("Configuration", "Reusable project files: regions, expected areas, settings", "Setup")
    picture(s, "10_configuration.png", Inches(0.45), AREA_TOP, Inches(9.2), AREA_H)
    bullets(s, Inches(9.85), Inches(1.6), Inches(3.3), Inches(5.3), [
        "Add / edit / remove / reorder",
        "End address or +size",
        "Validation: ranges, duplicates, overlaps",
        "Gap fill for sparse HEX vs full BIN (opt-in)",
        "Saved as JSON, used by GUI and CLI",
    ], size=14, space=10)
    notes(s, "Expected areas are never assumed — the engineer configures them explicitly.")


def s_controller():
    s = content_slide("Controller identification", "Exact part from symbols, family from the image layout", "Highlight")
    picture(s, "14_controller_pic32.png", Inches(0.45), AREA_TOP, Inches(8.6), AREA_H)
    bullets(s, Inches(9.25), Inches(1.6), Inches(3.85), Inches(5.3), [
        ("PIC32MX575F512L ", "from the XC32 map — 512 KB flash, 100 pins"),
        ("STM32F411RET ", "from ELF debug info — Cortex-M4F, 512 KB, LQFP64"),
        ("HEX / BIN only: ", "family from flash layout, config words, vector table"),
        ("Warning ", "if A and B target different controllers"),
    ], size=14, space=10)
    notes(s, "Every identification lists its evidence and confidence. An application at 0x08040000 is flagged as "
             "sitting behind a bootloader.")


def s_details():
    s = content_slide("Metadata · Checksums · ELF", "Supporting analysis", "Demo")
    cw = Inches(4.05)
    for i, (img, cap) in enumerate((("11_metadata.png", "Heuristic metadata: version, build date/time, git, "
                                                        "product, board"),
                                    ("12_checksums.png", "CRC-32 / CRC-16 / SHA-256 over file, image, region or "
                                                         "range"),
                                    ("13_elf.png", "ELF sections, build info and symbol changes"))):
        x = Inches(0.45) + i * (cw + Inches(0.2))
        picture(s, img, x, Inches(1.45), cw, Inches(4.3))
        text(s, x, Inches(5.9), cw, Inches(0.9), cap, size=13, color=MUTED, align=PP_ALIGN.CENTER)
    notes(s, "Metadata detection is heuristic and says so. Checksums never assume the firmware's own CRC algorithm.")


def s_cli():
    s = content_slide("Command line & CI", "Same engine, scriptable — exit codes gate a release", "Automation")
    picture(s, "16_cli.png", Inches(0.45), AREA_TOP, Inches(8.0), AREA_H, border=False)
    rect(s, Inches(8.75), Inches(1.55), Inches(4.2), Inches(2.35), WHITE, line=RGBColor(0xD8, 0xDE, 0xE6))
    text(s, Inches(8.95), Inches(1.65), Inches(3.9), Inches(0.4), "Exit codes", size=16, bold=True)
    bullets(s, Inches(8.95), Inches(2.1), Inches(3.9), Inches(1.8), [
        ("0 ", "identical or expected changes only"),
        ("1 ", "unexpected changes"),
        ("2 ", "error (missing / malformed file, bad config)"),
    ], size=13, space=4)
    bullets(s, Inches(8.75), Inches(4.15), Inches(4.2), Inches(2.8), [
        "--report / --csv / --json",
        "--symbols-a / --symbols-b",
        "--regions / --ignore",
        "--fill, --merge-gap, --checksum",
    ], size=13, space=5)
    notes(s, "Show the one-line CI usage: firmware-diff old.bin new.bin --regions project.json --report r.html. "
             "A non-zero exit code fails the pipeline on unexpected changes.")


def s_report():
    s = content_slide("HTML engineering report", "Self-contained, printable, shareable", "Output")
    picture(s, "17b_html_report_top.png", Inches(0.45), AREA_TOP, Inches(6.0), AREA_H)
    bullets(s, Inches(6.75), Inches(1.6), Inches(6.2), Inches(5.3), [
        "Firmware A / B information incl. controller",
        "SHA-256 hashes, overall summary, memory map",
        "Region comparison",
        "Changed blocks with function names",
        "Expected and unexpected changes",
        "Function and module changes, side effects",
        "Metadata differences, ELF comparison, timestamp",
        "CSV exports for blocks and functions; JSON summary",
    ], size=15, space=8)
    notes(s, "Open docs/sample_comparison_report.html in a browser to show the full report.")


def s_quality():
    s = content_slide("Quality and performance", "Verified on synthetic and production firmware", "Quality")
    stats = [("84", "automated tests", GREEN), ("0.1 s", "64 MB comparison", BLUE),
             ("2", "production firmware families", PURPLE), ("0", "runtime dependencies", SLATE)]
    cw = Inches(2.95)
    for i, (big, cap, c) in enumerate(stats):
        x = Inches(0.55) + i * (cw + Inches(0.17))
        rect(s, x, Inches(1.5), cw, Inches(1.65), WHITE, line=RGBColor(0xD8, 0xDE, 0xE6))
        text(s, x, Inches(1.6), cw, Inches(0.9), big, size=40, bold=True, color=c, align=PP_ALIGN.CENTER)
        text(s, x, Inches(2.5), cw, Inches(0.5), cap, size=14, color=MUTED, align=PP_ALIGN.CENTER)
    bullets(s, Inches(0.6), Inches(3.5), Inches(12.2), Inches(3.4), [
        ("Engine: ", "randomised tests against a naive byte-by-byte reference; interval algebra against Python sets."),
        ("Formats: ", "malformed / truncated HEX and ELF, overlapping records, files verified unmodified."),
        ("Functions: ", "real Thumb-2 and MIPS32 machine code, aliases, linker labels, renamed compiler symbols."),
        ("Production: ", "STM32F411 watchdog change → exactly 2 modified + 6 removed; PIC32MX575 release diff."),
        ("GUI: ", "smoke test drives the real window — load, compare, tabs, hex navigation, filters, search."),
    ], size=16, space=7)
    notes(s, "Mention the freeze: v1.0.0 baseline with a SHA-256 manifest of every file (docs/FREEZE_v1.0.0.md).")


def s_limits():
    s = content_slide("Limitations and next steps", "v1.0.0 is frozen — candidates for the next version", "Roadmap")
    bullets(s, Inches(0.6), Inches(1.5), Inches(6.0), Inches(5.3), [
        ("Heuristics: ", "metadata and controller identification are best-effort, with evidence shown"),
        ("Decoding: ", "MOVW/MOVT, microMIPS / MIPS16e and RISC-V not decoded yet"),
        ("Maps: ", "GNU ld and XC32 only (IAR / Keil: use their ELF)"),
        ("Formats: ", "no Motorola S-record yet"),
        ("Same-size MODIFIED ", "can be a side effect of header / struct changes"),
    ], size=16, space=10)
    rect(s, Inches(6.95), Inches(1.5), Inches(5.95), Inches(5.0), WHITE, line=RGBColor(0xD8, 0xDE, 0xE6))
    text(s, Inches(7.2), Inches(1.65), Inches(5.5), Inches(0.5), "Possible next steps", size=18, bold=True,
         color=BLUE)
    bullets(s, Inches(7.2), Inches(2.25), Inches(5.5), Inches(4.2), [
        "S-record and IAR / Keil map support",
        "RISC-V and MOVW/MOVT reference decoding",
        "Batch / folder comparison of release sets",
        "Git integration: compare two tags directly",
        "Packaged installer (PyInstaller scripts ready)",
    ], size=15, space=8)
    notes(s, "Invite feedback on priorities for the next version.")


def s_end():
    s = prs.slides.add_slide(BLANK)
    _page[0] += 1
    rect(s, 0, 0, W, H, SLATE)
    rect(s, 0, Inches(4.35), W, Inches(0.08), BLUE)
    text(s, Inches(0.8), Inches(2.1), Inches(11.5), Inches(1.0), "Thank you — questions?", size=40, bold=True,
         color=WHITE)
    text(s, Inches(0.8), Inches(3.15), Inches(11.5), Inches(0.8),
         "Byte-exact comparison · regions & expectations · real edits vs side effects · controller identification",
         size=17, color=RGBColor(0xCB, 0xD5, 0xE1))
    text(s, Inches(0.8), Inches(4.8), Inches(11.5), Inches(1.6), [
        "Run:  python firmware_compare_gui.py   ·   python firmware_diff.py old.bin new.bin --report r.html",
        "Docs:  docs/Firmware_Binary_Comparison_Tool_Design_v1.0.docx   ·   README.md   ·   docs/FREEZE_v1.0.0.md"],
         size=14, color=RGBColor(0x93, 0xC5, 0xFD))
    notes(s, "Close: where to find the tool, the design document and the freeze manifest.")


def main():
    for build in (s_title, s_problem, s_features, s_architecture, s_pipeline, s_files, s_summary, s_functions,
                  s_relocated, s_hex, s_regions, s_blocks_map, s_config, s_controller, s_details, s_cli, s_report,
                  s_quality, s_limits, s_end):
        build()
    prs.save(OUT)
    print("written", OUT, "slides:", len(prs.slides))


if __name__ == "__main__":
    main()
